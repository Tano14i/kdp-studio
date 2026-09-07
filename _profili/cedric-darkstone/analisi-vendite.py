#!/usr/bin/env python3
"""Estrae i numeri del catalogo Cedric Darkstone dal report KDP Orders.

I numeri di `vendite-2026.md` NON si scrivono a mano: si generano da qui.
Uso:
    python3 _profili/cedric-darkstone/analisi-vendite.py <KDP_Orders.xlsx> \
        > _profili/cedric-darkstone/vendite-2026.md

Esce con codice != 0 se un controllo non passa (vedi sezione CONTROLLI in
fondo all'output). Un avviso che nessuno legge e' peggio di nessun avviso.
"""
import sys, collections, unicodedata

# --- I cambi stanno qui, in un punto solo. Si cambiano qui. ------------------
# Indicativi, medi 2026. Servono a mettere in fila mercati diversi, non a fare
# contabilita': il dato vero e' la royalty per valuta, che l'output riporta
# sempre non convertita accanto alla stima in euro.
FX = {'EUR':1.00,'USD':0.92,'GBP':1.17,'AUD':0.60,'JPY':0.0062,
      'CAD':0.68,'INR':0.011,'MXN':0.05,'BRL':0.17,'PLN':0.23,'SEK':0.088}

# Royalty attesa per tipo di transazione: uno scostamento e' un guasto, non
# un dettaglio. KDP paga 60% sul cartaceo standard, 40% in distribuzione
# estesa, 70% o 35% sull'ebook a seconda della fascia di prezzo.
ATTESA = {'Standard - Paperback':{'60%'},
          'Expanded Distribution Channels':{'40%'},
          'Standard':{'70%','35%'},
          'Free - Promotion':{'70%','35%'}}

def breve(t):
    for chiave, nome in (('Diario','Diario del DM (IT)'),
                         ('First-Time','First-Time DM (EN)'),
                         ('Diventa','Diventa Master (IT)'),
                         ('500','500 Encounter Tables (EN)')):
        if t.startswith(chiave):
            return nome
    return t[:28]

def main(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    guasti = []

    def righe(foglio):
        ws = wb[foglio]
        hdr = [c.value for c in ws[1]]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            yield dict(zip(hdr, r))

    vendite = list(righe('Combined Sales'))
    kenp    = list(righe('KENP Read'))
    sommario= list(righe('Summary'))

    # ---- royalty per valuta, dato grezzo -----------------------------------
    per_valuta = collections.defaultdict(float)
    for d in vendite:
        c = d['Currency']
        if c not in FX:
            guasti.append(f"valuta senza cambio in tabella FX: {c}")
        per_valuta[c] += d['Royalty'] or 0.0

    # il Sommario include anche le pagine lette in Kindle Unlimited
    sommario_valuta = collections.defaultdict(float)
    for d in sommario:
        for k, v in d.items():
            if k.startswith('Royalty (') and v:
                sommario_valuta[k[9:-1]] += v

    tot_eur = sum(v * FX.get(c, 0) for c, v in sommario_valuta.items())
    ku_eur  = tot_eur - sum(v * FX.get(c, 0) for c, v in per_valuta.items())

    # ---- aggregati ---------------------------------------------------------
    T = lambda: {'eur':0.0,'eb':0,'pb':0,'free':0}
    titolo   = collections.defaultdict(T)
    mercato  = collections.defaultdict(lambda:{'eur':0.0,'u':0})
    formato  = collections.defaultdict(lambda:{'eur':0.0,'u':0})
    mese     = collections.defaultdict(float)
    tit_mese = collections.defaultdict(float)
    sku      = {}
    nomi     = collections.defaultdict(set)

    for d in vendite:
        t   = breve(d['Title'])
        eur = (d['Royalty'] or 0.0) * FX.get(d['Currency'], 0)
        u   = int(d['Net Units Sold'] or 0)
        tt  = d['Transaction Type']
        gratis = 'Free' in tt
        cart   = 'Paperback' in tt or 'Expanded' in tt
        f = 'Cartaceo' if cart else 'eBook'

        titolo[t]['eur'] += eur
        if gratis:
            titolo[t]['free'] += u
        else:
            titolo[t]['pb' if cart else 'eb'] += u
            mercato[d['Marketplace']]['u'] += u
            formato[f]['u'] += u
        mercato[d['Marketplace']]['eur'] += eur
        formato[f]['eur'] += eur
        mese[d['Royalty Date']] += eur
        tit_mese[(t, d['Royalty Date'])] += eur

        nomi[unicodedata.normalize('NFKC', d['Author Name'])].add(t)

        if not gratis:
            k = (t, f, d['Marketplace'], d['Currency'])
            # la royalty della riga copre `u` copie: qui serve quella per copia
            sku[k] = (d['Avg. List Price without tax'],
                      d['Avg. Delivery/Manufacturing cost'],
                      (d['Royalty'] or 0.0) / u if u else 0.0,
                      d['Royalty Type'])
        # CONTROLLI di riga
        att = ATTESA.get(tt)
        if att and d['Royalty Type'] not in att:
            guasti.append(f"{t} · {d['Marketplace']} · {tt}: royalty "
                          f"{d['Royalty Type']}, attesa {'/'.join(sorted(att))}")
        if not gratis and (d['Royalty'] or 0) == 0:
            guasti.append(f"{t} · {d['Marketplace']} · {d['Royalty Date']}: "
                          f"{u} copie vendute a {d['Avg. List Price without tax']} "
                          f"{d['Currency']} con royalty ZERO")

    if len(nomi) > 1:
        guasti.append("nome autore non identico su tutti i titoli: "
                      + " · ".join(f"{n!r} → {', '.join(sorted(v))}"
                                   for n, v in nomi.items()))

    kenp_tit = collections.defaultdict(float)
    for d in kenp:
        kenp_tit[breve(d['Title'])] += d['Kindle Edition Normalized Page (KENP) Read'] or 0

    # ---- stampa ------------------------------------------------------------
    P = print
    P("<!-- GENERATO da analisi-vendite.py — non modificare a mano. -->")
    P(f"# Vendite Cedric Darkstone — {min(mese)} → {max(mese)}\n")
    P(f"Royalty totale stimata: **{tot_eur:.2f} EUR** "
      f"in {len(mese)} mesi = **{tot_eur/len(mese):.2f} EUR/mese**.  ")
    P(f"Di cui Kindle Unlimited: {ku_eur:.2f} EUR "
      f"({100*ku_eur/tot_eur:.1f}%). Cambi indicativi, vedi FX nello script.\n")

    P("## Royalty per valuta (dato grezzo, non convertito)\n")
    P("| Valuta | Vendite | Totale col KU |")
    P("|---|---:|---:|")
    for c in sorted(sommario_valuta, key=lambda c: -sommario_valuta[c] * FX.get(c, 0)):
        P(f"| {c} | {per_valuta.get(c,0):.2f} | {sommario_valuta[c]:.2f} |")

    P("\n## Per titolo\n")
    P("| Titolo | EUR | quota | eBook | Cartaceo | Gratis | EUR/copia | KENP |")
    P("|---|---:|---:|---:|---:|---:|---:|---:|")
    for t, d in sorted(titolo.items(), key=lambda x: -x[1]['eur']):
        u = d['eb'] + d['pb']
        P(f"| {t} | {d['eur']:.2f} | {100*d['eur']/tot_eur:.0f}% | {d['eb']} | "
          f"{d['pb']} | {d['free']} | {d['eur']/u if u else 0:.2f} | "
          f"{kenp_tit.get(t,0):.0f} |")

    P("\n## Per formato\n")
    P("| Formato | EUR | quota | Copie | EUR/copia |")
    P("|---|---:|---:|---:|---:|")
    for f, d in sorted(formato.items(), key=lambda x: -x[1]['eur']):
        P(f"| {f} | {d['eur']:.2f} | {100*d['eur']/tot_eur:.0f}% | {d['u']} | "
          f"{d['eur']/d['u'] if d['u'] else 0:.2f} |")

    P("\n## Per mercato\n")
    P("| Mercato | EUR | quota | Copie |")
    P("|---|---:|---:|---:|")
    for m, d in sorted(mercato.items(), key=lambda x: -x[1]['eur']):
        P(f"| {m} | {d['eur']:.2f} | {100*d['eur']/tot_eur:.0f}% | {d['u']} |")

    P("\n## Per mese, e chi lo ha prodotto\n")
    mesi = sorted(mese)
    ordine = sorted(titolo, key=lambda t: -titolo[t]['eur'])
    P("| Mese | " + " | ".join(ordine) + " | Totale |")
    P("|---" * (len(ordine) + 2) + "|")
    for m in mesi:
        P(f"| {m} | " + " | ".join(f"{tit_mese.get((t,m),0):.2f}" for t in ordine)
          + f" | **{mese[m]:.2f}** |")

    P(f"\nL'ultimo mese ({max(mesi)}) e' parziale: il report si ferma al "
      "giorno dello scarico.\n")

    P("\n## Economia per copia, SKU per SKU\n")
    P("| Titolo | Formato | Mercato | Prezzo | Costo | Royalty | Tasso | Margine |")
    P("|---|---|---|---:|---:|---:|---:|---:|")
    for (t, f, m, c), (lp, dc, roy, rt) in sorted(sku.items()):
        cost = f"{dc:.2f}" if isinstance(dc, (int, float)) else str(dc)
        P(f"| {t} | {f} | {m} | {lp:.2f} {c} | {cost} | {roy:.2f} | {rt} | "
          f"{100*roy/lp:.0f}% |")

    P("\n## CONTROLLI\n")
    P(f"Righe vendita lette: {len(vendite)} · titoli: {len(titolo)} · "
      f"mercati: {len(mercato)} · mesi: {len(mese)}.\n")
    guasti = list(dict.fromkeys(guasti))
    if guasti:
        P(f"**{len(guasti)} da guardare:**\n")
        for g in guasti:
            P(f"- {g}")
    else:
        P("Nessun guasto.")
    return 1 if guasti else 0

if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
